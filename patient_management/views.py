import base64
import os
from io import BytesIO

import requests

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.http import HttpResponse
from django.utils import timezone
from django.shortcuts import get_object_or_404, redirect, render

from openpyxl import Workbook
from openpyxl import load_workbook

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from .forms import PatientForm
from .models import Patient


@login_required
def patient_list(request):

    search = request.GET.get(
        'search',
        ''
    ).strip()

    patients = Patient.objects.select_related(
        'user'
    ).order_by(
        '-created_at'
    )

    if search:

        patients = patients.filter(
            patient_id__icontains=search
        )

    paginator = Paginator(
        patients,
        10
    )

    page_number = request.GET.get(
        'page'
    )

    patients = paginator.get_page(
        page_number
    )

    # AJAX request:
    # Return only the patient table section.
    if request.headers.get('X-Requested-With') == 'XMLHttpRequest':

        return render(
            request,
            'patients/partials/patient_table.html',
            {
                'patients': patients,
                'search': search,
            }
        )

    return render(
        request,
        'patients/patient_list.html',
        {
            'patients': patients,
            'search': search,
        }
    )

@login_required
def patient_create(request):

    if request.method == 'POST':

        form = PatientForm(
            request.POST,
            request.FILES
        )

        if form.is_valid():

            form.save()

            return redirect(
                'patient_list'
            )

    else:

        form = PatientForm()

    return render(
        request,
        'patients/patient_form.html',
        {
            'form': form
        }
    )

@login_required
def patient_edit(request, patient_id):

    patient = get_object_or_404(
        Patient,
        id=patient_id
    )

    # =================================================
    # AJAX EDIT REQUEST
    # =================================================

    if request.headers.get('X-Requested-With') == 'XMLHttpRequest':

        if request.method == 'POST':

            form = PatientForm(
                request.POST,
                request.FILES,
                instance=patient
            )

            if form.is_valid():

                form.save()

                return HttpResponse(
                    'success'
                )

        else:

            form = PatientForm(
                instance=patient
            )

        return render(
            request,
            'patients/partials/patient_edit_form.html',
            {
                'form': form,
                'patient': patient,
            }
        )

    # =================================================
    # NORMAL EDIT REQUEST
    # =================================================

    if request.method == 'POST':

        form = PatientForm(
            request.POST,
            request.FILES,
            instance=patient
        )

        if form.is_valid():

            form.save()

            return redirect(
                'patient_list'
            )

    else:

        form = PatientForm(
            instance=patient
        )

    return render(
        request,
        'patients/patient_form.html',
        {
            'form': form,
            'edit_mode': True
        }
    )
@login_required
def patient_delete(request, patient_id):

    patient = get_object_or_404(
        Patient,
        id=patient_id
    )

    # =================================================
    # AJAX DELETE REQUEST
    # =================================================

    if request.headers.get(
        'X-Requested-With'
    ) == 'XMLHttpRequest':

        # Show confirmation modal
        if request.method == 'GET':

            return render(
                request,
                'patients/patient_delete.html',
                {
                    'patient': patient
                }
            )

        # Confirm deletion
        if request.method == 'POST':

            patient.delete()

            return HttpResponse(
                'success'
            )


    # =================================================
    # NORMAL DELETE REQUEST
    # =================================================

    if request.method == 'POST':

        patient.delete()

        return redirect(
            'patient_list'
        )


    return render(
        request,
        'patients/patient_delete.html',
        {
            'patient': patient
        }
    )

@login_required
def patient_upload_excel(request):

    if request.method == 'POST':

        excel_file = request.FILES.get(
            'excel_file'
        )

        if not excel_file:

            messages.error(
                request,
                'Please select an Excel file.'
            )

            return redirect(
                'patient_list'
            )

        if not excel_file.name.endswith('.xlsx'):

            messages.error(
                request,
                'Only .xlsx Excel files are supported.'
            )

            return redirect(
                'patient_list'
            )

        try:

            workbook = load_workbook(
                excel_file
            )

            sheet = workbook.active

            created_count = 0
            skipped_count = 0

            for row in sheet.iter_rows(
                min_row=2,
                values_only=True
            ):

                if not row[0]:
                    continue

                patient_id = str(
                    row[0]
                ).strip()

                if Patient.objects.filter(
                    patient_id=patient_id
                ).exists():

                    skipped_count += 1

                    continue

                Patient.objects.create(

                    patient_id=patient_id,

                    date_of_birth=row[1],

                    gender=str(
                        row[2]
                    ).strip().lower(),

                    phone=str(
                        row[3]
                    ).strip(),

                    address=str(
                        row[4]
                    ).strip(),

                    blood_group=(
                        str(row[5]).strip()
                        if row[5]
                        else ''
                    ),

                    emergency_contact=(
                        str(row[6]).strip()
                        if row[6]
                        else ''
                    ),
                )

                created_count += 1

            messages.success(
                request,
                f'{created_count} patients imported successfully. '
                f'{skipped_count} duplicate patients skipped.'
            )

        except Exception as error:

            print(
                'Excel upload error:',
                error
            )

            messages.error(
                request,
                'Unable to process the Excel file. '
                'Please check the file format and data.'
            )

        return redirect(
            'patient_list'
        )

    return redirect(
        'patient_list'
    )


@login_required
def patient_download_excel(request):

    patients = Patient.objects.order_by(
        'patient_id'
    )

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = 'Patients'

    headers = [
        'Patient ID',
        'Date of Birth',
        'Gender',
        'Phone',
        'Address',
        'Blood Group',
        'Emergency Contact',
    ]

    worksheet.append(headers)

    for patient in patients:

        worksheet.append([
            patient.patient_id,
            patient.date_of_birth,
            patient.get_gender_display(),
            patient.phone,
            patient.address,
            patient.blood_group or '',
            patient.emergency_contact or '',
        ])

    column_widths = {
        'A': 15,
        'B': 18,
        'C': 15,
        'D': 15,
        'E': 35,
        'F': 15,
        'G': 20,
    }

    for column, width in column_widths.items():

        worksheet.column_dimensions[
            column
        ].width = width

    excel_buffer = BytesIO()

    workbook.save(
        excel_buffer
    )

    excel_bytes = excel_buffer.getvalue()

    try:

        api_key = os.environ.get(
            'NYLAS_API_KEY'
        )

        grant_id = os.environ.get(
            'NYLAS_GRANT_ID'
        )

        if not api_key or not grant_id:

            raise Exception(
                'Nylas configuration is missing.'
            )

        encoded_file = base64.b64encode(
            excel_bytes
        ).decode('utf-8')

        url = (
            'https://api.us.nylas.com/v3/grants/'
            f'{grant_id}/messages/send'
        )

        headers = {
            'Authorization': f'Bearer {api_key}',
            'Content-Type': 'application/json',
        }

        data = {
            'to': [
                {
                    'email': 'muthyalatejas8055@gmail.com',
                }
            ],
            'subject': (
                'MediCare Plus - '
                'Patient Records Excel Report'
            ),
            'body': (
                '<h2>MediCare Plus Hospital</h2>'
                '<p>Patient records Excel report '
                'generated successfully.</p>'
                '<p>The Excel file is attached to '
                'this email.</p>'
            ),
            'attachments': [
                {
                    'filename': (
                        'medicare_plus_patients.xlsx'
                    ),
                    'content': encoded_file,
                    'content_type': (
                        'application/vnd.openxmlformats-officedocument'
                        '.spreadsheetml.sheet'
                    ),
                }
            ],
        }

        response = requests.post(
            url,
            headers=headers,
            json=data,
            timeout=30
        )

        print(
            'Nylas Excel email status:',
            response.status_code
        )

        print(
            'Nylas Excel email response:',
            response.text
        )

        response.raise_for_status()

        messages.success(
            request,
            'Patient Excel report was sent successfully by email.'
        )

    except Exception as error:

        print(
            'Nylas Excel email error:',
            error
        )

        messages.warning(
            request,
            'Excel file downloaded, but the email could not be sent.'
        )

    response = HttpResponse(
        excel_bytes,
        content_type=(
            'application/vnd.openxmlformats-officedocument'
            '.spreadsheetml.sheet'
        )
    )

    response['Content-Disposition'] = (
        'attachment; '
        'filename="medicare_plus_patients.xlsx"'
    )

    return response

@login_required
def patient_download_pdf(request, patient_id):

    patient = get_object_or_404(
        Patient,
        id=patient_id
    )

    pdf_buffer = BytesIO()

    pdf = canvas.Canvas(
        pdf_buffer,
        pagesize=A4
    )

    width, height = A4

    # =================================================
    # PAGE SETTINGS
    # =================================================

    left = 45
    right = width - 45

    # =================================================
    # HEADER
    # =================================================

    pdf.setFillColorRGB(
        0.08,
        0.32,
        0.45
    )

    pdf.rect(
        0,
        height - 105,
        width,
        105,
        fill=1,
        stroke=0
    )

    pdf.setFillColorRGB(
        1,
        1,
        1
    )

    pdf.setFont(
        'Helvetica-Bold',
        23
    )

    pdf.drawString(
        left,
        height - 45,
        'MEDICARE PLUS'
    )

    pdf.setFont(
        'Helvetica',
        10
    )

    pdf.drawString(
        left,
        height - 63,
        'HOSPITAL MANAGEMENT SYSTEM'
    )

    pdf.setFont(
        'Helvetica',
        9
    )

    pdf.drawRightString(
        right,
        height - 45,
        'PATIENT SERVICES'
    )

    pdf.drawRightString(
        right,
        height - 61,
        'Medical Records Department'
    )

    # =================================================
    # REPORT TITLE
    # =================================================

    pdf.setFillColorRGB(
        0.08,
        0.32,
        0.45
    )

    pdf.setFont(
        'Helvetica-Bold',
        18
    )

    pdf.drawString(
        left,
        height - 140,
        'PATIENT MEDICAL REPORT'
    )

    pdf.setFont(
        'Helvetica',
        9
    )

    pdf.setFillColorRGB(
        0.35,
        0.35,
        0.35
    )

    pdf.drawString(
        left,
        height - 157,
        'Confidential Patient Record'
    )

    # =================================================
    # PATIENT ID BADGE
    # =================================================

    badge_width = 135
    badge_height = 32
    badge_x = right - badge_width
    badge_y = height - 153

    pdf.setFillColorRGB(
        0.90,
        0.95,
        0.97
    )

    pdf.roundRect(
        badge_x,
        badge_y,
        badge_width,
        badge_height,
        6,
        fill=1,
        stroke=0
    )

    pdf.setFillColorRGB(
        0.08,
        0.32,
        0.45
    )

    pdf.setFont(
        'Helvetica-Bold',
        9
    )

    pdf.drawString(
        badge_x + 10,
        badge_y + 19,
        'PATIENT ID'
    )

    pdf.setFont(
        'Helvetica-Bold',
        12
    )

    pdf.drawString(
        badge_x + 10,
        badge_y + 7,
        str(patient.patient_id)
    )

    # =================================================
    # PATIENT INFORMATION
    # =================================================

    section_y = height - 200

    pdf.setFillColorRGB(
        0.08,
        0.32,
        0.45
    )

    pdf.setFont(
        'Helvetica-Bold',
        12
    )

    pdf.drawString(
        left,
        section_y,
        'PATIENT INFORMATION'
    )

    pdf.setLineWidth(
        1
    )

    pdf.line(
        left,
        section_y - 7,
        right,
        section_y - 7
    )

    # =================================================
    # INFORMATION TABLE
    # =================================================

    table_top = section_y - 22
    row_height = 34

    table_rows = [
        (
            'Date of Birth',
            str(patient.date_of_birth),
            'Gender',
            patient.get_gender_display()
        ),
        (
            'Phone',
            patient.phone,
            'Blood Group',
            patient.blood_group or 'Not provided'
        ),
        (
            'Emergency Contact',
            patient.emergency_contact or 'Not provided',
            '',
            ''
        ),
    ]

    table_height = row_height * len(
        table_rows
    )

    pdf.setStrokeColorRGB(
        0.80,
        0.84,
        0.86
    )

    pdf.setFillColorRGB(
        0.97,
        0.98,
        0.98
    )

    pdf.rect(
        left,
        table_top - table_height,
        right - left,
        table_height,
        fill=1,
        stroke=1
    )

    # Vertical divider
    middle = width / 2

    pdf.line(
        middle,
        table_top,
        middle,
        table_top - table_height
    )

    for index in range(
        1,
        len(table_rows)
    ):

        y = table_top - (
            index * row_height
        )

        pdf.line(
            left,
            y,
            right,
            y
        )

    for index, row in enumerate(
        table_rows
    ):

        y = table_top - (
            index * row_height
        ) - 14

        label1 = row[0]
        value1 = row[1]
        label2 = row[2]
        value2 = row[3]

        pdf.setFont(
            'Helvetica-Bold',
            8.5
        )

        pdf.setFillColorRGB(
            0.35,
            0.35,
            0.35
        )

        pdf.drawString(
            left + 12,
            y,
            label1
        )

        pdf.setFont(
            'Helvetica',
            9.5
        )

        pdf.setFillColorRGB(
            0.10,
            0.10,
            0.10
        )

        pdf.drawString(
            left + 12,
            y - 13,
            str(value1)
        )

        if label2:

            pdf.setFont(
                'Helvetica-Bold',
                8.5
            )

            pdf.setFillColorRGB(
                0.35,
                0.35,
                0.35
            )

            pdf.drawString(
                middle + 12,
                y,
                label2
            )

            pdf.setFont(
                'Helvetica',
                9.5
            )

            pdf.setFillColorRGB(
                0.10,
                0.10,
                0.10
            )

            pdf.drawString(
                middle + 12,
                y - 13,
                str(value2)
            )

    # =================================================
    # ADDRESS
    # =================================================

    address_section_y = (
        table_top
        - table_height
        - 35
    )

    pdf.setFillColorRGB(
        0.08,
        0.32,
        0.45
    )

    pdf.setFont(
        'Helvetica-Bold',
        12
    )

    pdf.drawString(
        left,
        address_section_y,
        'ADDRESS'
    )

    pdf.setLineWidth(
        1
    )

    pdf.line(
        left,
        address_section_y - 7,
        right,
        address_section_y - 7
    )

    address_box_top = (
        address_section_y - 20
    )

    address_box_height = 58

    pdf.setFillColorRGB(
        0.97,
        0.98,
        0.98
    )

    pdf.setStrokeColorRGB(
        0.80,
        0.84,
        0.86
    )

    pdf.roundRect(
        left,
        address_box_top - address_box_height,
        right - left,
        address_box_height,
        5,
        fill=1,
        stroke=1
    )

    pdf.setFillColorRGB(
        0.10,
        0.10,
        0.10
    )

    pdf.setFont(
        'Helvetica',
        10
    )

    address = str(
        patient.address or 'Not provided'
    )

    max_chars = 85

    address_lines = []

    remaining = address

    while remaining:

        if len(remaining) <= max_chars:

            address_lines.append(
                remaining
            )

            break

        split_position = (
            remaining.rfind(
                ' ',
                0,
                max_chars
            )
        )

        if split_position == -1:

            split_position = max_chars

        address_lines.append(
            remaining[:split_position]
        )

        remaining = remaining[
            split_position:
        ].strip()

    address_y = (
        address_box_top - 21
    )

    for line in address_lines[:2]:

        pdf.drawString(
            left + 12,
            address_y,
            line
        )

        address_y -= 15

    # =================================================
    # MEDICAL DOCUMENT
    # =================================================

    medical_y = (
        address_box_top
        - address_box_height
        - 35
    )

    pdf.setFillColorRGB(
        0.08,
        0.32,
        0.45
    )

    pdf.setFont(
        'Helvetica-Bold',
        12
    )

    pdf.drawString(
        left,
        medical_y,
        'MEDICAL DOCUMENT'
    )

    pdf.line(
        left,
        medical_y - 7,
        right,
        medical_y - 7
    )

    document_box_top = (
        medical_y - 20
    )

    document_box_height = 55

    pdf.setFillColorRGB(
        0.97,
        0.98,
        0.98
    )

    pdf.setStrokeColorRGB(
        0.80,
        0.84,
        0.86
    )

    pdf.roundRect(
        left,
        document_box_top - document_box_height,
        right - left,
        document_box_height,
        5,
        fill=1,
        stroke=1
    )

    pdf.setFillColorRGB(
        0.10,
        0.10,
        0.10
    )

    pdf.setFont(
        'Helvetica-Bold',
        9
    )

    pdf.drawString(
        left + 12,
        document_box_top - 20,
        'Medical Report File'
    )

    pdf.setFont(
        'Helvetica',
        10
    )

    if patient.medical_report:

        pdf.drawString(
            left + 150,
            document_box_top - 20,
            str(
                patient.medical_report.name
            )
        )

        pdf.setFont(
            'Helvetica',
            8.5
        )

        pdf.drawString(
            left + 150,
            document_box_top - 36,
            'File available in patient record'
        )

    else:

        pdf.drawString(
            left + 150,
            document_box_top - 20,
            'Not provided'
        )

    # =================================================
    # REPORT DETAILS
    # =================================================

    report_y = (
        document_box_top
        - document_box_height
        - 35
    )

    pdf.setFillColorRGB(
        0.08,
        0.32,
        0.45
    )

    pdf.setFont(
        'Helvetica-Bold',
        12
    )

    pdf.drawString(
        left,
        report_y,
        'REPORT DETAILS'
    )

    pdf.line(
        left,
        report_y - 7,
        right,
        report_y - 7
    )

    generated_time = (
        timezone.localtime()
        .strftime(
            '%d %B %Y, %I:%M %p'
        )
    )

    generated_by = (
        request.user.email
        if request.user.email
        else request.user.username
    )

    detail_y = report_y - 27

    pdf.setFillColorRGB(
        0.10,
        0.10,
        0.10
    )

    pdf.setFont(
        'Helvetica-Bold',
        9
    )

    pdf.drawString(
        left,
        detail_y,
        'Generated On'
    )

    pdf.setFont(
        'Helvetica',
        9
    )

    pdf.drawString(
        left + 110,
        detail_y,
        generated_time
    )

    pdf.setFont(
        'Helvetica-Bold',
        9
    )

    pdf.drawString(
        left,
        detail_y - 18,
        'Generated By'
    )

    pdf.setFont(
        'Helvetica',
        9
    )

    pdf.drawString(
        left + 110,
        detail_y - 18,
        generated_by
    )

    # =================================================
    # FOOTER
    # =================================================

    pdf.setStrokeColorRGB(
        0.80,
        0.84,
        0.86
    )

    pdf.line(
        left,
        60,
        right,
        60
    )

    pdf.setFillColorRGB(
        0.35,
        0.35,
        0.35
    )

    pdf.setFont(
        'Helvetica',
        8
    )

    pdf.drawString(
        left,
        45,
        'MediCare Plus Hospital'
    )

    pdf.drawRightString(
        right,
        45,
        'Patient Medical Record'
    )

    pdf.drawCentredString(
        width / 2,
        28,
        'This document is generated electronically.'
    )

    # =================================================
    # FINISH PDF
    # =================================================

    pdf.save()

    pdf_bytes = pdf_buffer.getvalue()

    # =================================================
    # SEND PDF THROUGH NYLAS
    # =================================================

    try:

        api_key = os.environ.get(
            'NYLAS_API_KEY'
        )

        grant_id = os.environ.get(
            'NYLAS_GRANT_ID'
        )

        if not api_key or not grant_id:

            raise Exception(
                'Nylas configuration is missing.'
            )

        encoded_file = base64.b64encode(
            pdf_bytes
        ).decode('utf-8')

        url = (
            'https://api.us.nylas.com/v3/grants/'
            f'{grant_id}/messages/send'
        )

        headers = {
            'Authorization': (
                f'Bearer {api_key}'
            ),
            'Content-Type': 'application/json',
        }

        data = {
            'to': [
                {
                    'email': (
                        'muthyalatejas8055@gmail.com'
                    ),
                }
            ],
            'subject': (
                'MediCare Plus - '
                f'Patient Report - '
                f'{patient.patient_id}'
            ),
            'body': (
                '<h2>MediCare Plus Hospital</h2>'
                '<p>Your professionally formatted '
                'patient medical report has been '
                'generated successfully.</p>'
                '<p>The PDF report is attached '
                'to this email.</p>'
            ),
            'attachments': [
                {
                    'filename': (
                        f'patient_'
                        f'{patient.patient_id}.pdf'
                    ),
                    'content': encoded_file,
                    'content_type': (
                        'application/pdf'
                    ),
                }
            ],
        }

        response = requests.post(
            url,
            headers=headers,
            json=data,
            timeout=30
        )

        print(
            'Nylas email status:',
            response.status_code
        )

        print(
            'Nylas email response:',
            response.text
        )

        response.raise_for_status()

        messages.success(
            request,
            'Patient PDF report was sent successfully by email.'
        )

    except Exception as error:

        print(
            'Nylas email error:',
            error
        )

        messages.warning(
            request,
            'PDF downloaded, but the email could not be sent.'
        )

    # =================================================
    # DOWNLOAD PDF
    # =================================================

    response = HttpResponse(
        pdf_bytes,
        content_type='application/pdf'
    )

    response['Content-Disposition'] = (
        f'attachment; '
        f'filename="patient_'
        f'{patient.patient_id}.pdf"'
    )

    return response