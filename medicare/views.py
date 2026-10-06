from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.cache import never_cache

from openpyxl import Workbook, load_workbook

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from appointment_management.models import Appointment
from doctor_management.models import Doctor
from patient_management.models import Patient


@login_required
@never_cache
def home(request):
    """
    Main MediCare Plus dashboard.
    Shows statistics, patient records and recent appointments.
    """

    # Patient pagination - 10 records per page
    patients = Patient.objects.order_by('-created_at')

    paginator = Paginator(patients, 10)
    page_number = request.GET.get('page')
    patients_page = paginator.get_page(page_number)

    context = {
        'patient_count': Patient.objects.count(),
        'doctor_count': Doctor.objects.count(),
        'appointment_count': Appointment.objects.count(),

        'pending_count': Appointment.objects.filter(
            status='pending'
        ).count(),

        'confirmed_count': Appointment.objects.filter(
            status='confirmed'
        ).count(),

        'completed_count': Appointment.objects.filter(
            status='completed'
        ).count(),

        'cancelled_count': Appointment.objects.filter(
            status='cancelled'
        ).count(),

        'patients': patients_page,

        'recent_appointments': Appointment.objects.select_related(
            'patient',
            'doctor'
        ).order_by(
            '-appointment_date',
            '-appointment_time'
        )[:5],
    }

    return render(request, 'home.html', context)


@login_required
def dashboard_upload_pdf(request):
    """
    Upload a PDF medical report for a patient.
    """

    if request.method != 'POST':
        return redirect('home')

    patient_id = request.POST.get('patient_id', '').strip()
    pdf_file = request.FILES.get('pdf_file')

    if not patient_id:
        messages.error(request, 'Please enter a Patient ID.')
        return redirect('home')

    if not pdf_file:
        messages.error(request, 'Please select a PDF file.')
        return redirect('home')

    if not pdf_file.name.lower().endswith('.pdf'):
        messages.error(request, 'Only PDF files are allowed.')
        return redirect('home')

    if pdf_file.size > 5 * 1024 * 1024:
        messages.error(request, 'PDF file size must be 5 MB or less.')
        return redirect('home')

    try:
        patient = Patient.objects.get(patient_id=patient_id)

    except Patient.DoesNotExist:
        messages.error(
            request,
            f'Patient ID "{patient_id}" was not found.'
        )
        return redirect('home')

    patient.medical_report = pdf_file
    patient.save()

    messages.success(
        request,
        f'PDF medical report uploaded successfully for {patient.patient_id}.'
    )

    return redirect('home')


@login_required
def dashboard_upload_excel(request):
    """
    Upload patient records from an Excel file.
    """

    if request.method != 'POST':
        return redirect('home')

    excel_file = request.FILES.get('excel_file')

    if not excel_file:
        messages.error(request, 'Please select an Excel file.')
        return redirect('home')

    if not excel_file.name.lower().endswith('.xlsx'):
        messages.error(request, 'Only .xlsx Excel files are supported.')
        return redirect('home')

    try:
        workbook = load_workbook(excel_file)
        worksheet = workbook.active

        created_count = 0
        skipped_count = 0

        for row in worksheet.iter_rows(
            min_row=2,
            values_only=True
        ):

            if not row[0]:
                continue

            patient_id = str(row[0]).strip()

            if Patient.objects.filter(
                patient_id=patient_id
            ).exists():
                skipped_count += 1
                continue

            if not row[1]:
                skipped_count += 1
                continue

            gender = str(row[2]).strip().lower()

            if gender not in ['male', 'female', 'other']:
                skipped_count += 1
                continue

            Patient.objects.create(
                patient_id=patient_id,
                date_of_birth=row[1],
                gender=gender,
                phone=str(row[3]).strip() if row[3] else '',
                address=str(row[4]).strip() if row[4] else '',
                blood_group=str(row[5]).strip() if row[5] else '',
                emergency_contact=(
                    str(row[6]).strip()
                    if row[6]
                    else ''
                ),
            )

            created_count += 1

        messages.success(
            request,
            f'{created_count} patients imported successfully. '
            f'{skipped_count} records skipped.'
        )

    except Exception:
        messages.error(
            request,
            'Unable to process the Excel file. '
            'Please check the Excel format and data.'
        )

    return redirect('home')

@login_required
def dashboard_download_pdf(request):
    """
    Generate a patient's PDF report, email it to the
    MediCare Plus email address through Nylas,
    and download it in the browser.
    """

    import base64
    import os
    import requests

    from io import BytesIO
    from django.utils import timezone

    patient_id = request.GET.get('patient_id', '').strip()

    if not patient_id:
        messages.error(
            request,
            'Please enter a Patient ID to download the PDF.'
        )
        return redirect('home')

    patient = get_object_or_404(
        Patient,
        patient_id=patient_id
    )

    # Create PDF in memory
    pdf_buffer = BytesIO()

    pdf = canvas.Canvas(
        pdf_buffer,
        pagesize=A4
    )

    width, height = A4

    y = height - 60

    pdf.setFont(
        'Helvetica-Bold',
        22
    )

    pdf.drawString(
        50,
        y,
        'MediCare Plus Hospital'
    )

    y -= 40

    pdf.setFont(
        'Helvetica-Bold',
        16
    )

    pdf.drawString(
        50,
        y,
        'Patient Medical Report'
    )

    y -= 35

    pdf.setFont(
        'Helvetica',
        11
    )

    patient_details = [
        f'Patient ID: {patient.patient_id}',
        f'Date of Birth: {patient.date_of_birth}',
        f'Gender: {patient.get_gender_display()}',
        f'Phone: {patient.phone}',
        f'Blood Group: {patient.blood_group or "Not provided"}',
        f'Emergency Contact: '
        f'{patient.emergency_contact or "Not provided"}',
        f'Address: {patient.address}',
    ]

    for detail in patient_details:
        pdf.drawString(
            50,
            y,
            detail
        )
        y -= 25

    y -= 15

    pdf.setFont(
        'Helvetica-Oblique',
        10
    )

    pdf.drawString(
        50,
        y,
        'Generated by MediCare Plus Hospital Management System'
    )

    pdf.save()

    pdf_data = pdf_buffer.getvalue()
    pdf_buffer.close()

    filename = f'patient_{patient.patient_id}.pdf'

    # Send PDF through Nylas
    try:
        nylas_api_key = os.environ.get('NYLAS_API_KEY')
        nylas_grant_id = os.environ.get('NYLAS_GRANT_ID')

        if not nylas_api_key or not nylas_grant_id:
            raise ValueError(
                'Nylas API key or Grant ID is not configured.'
            )

        encoded_pdf = base64.b64encode(
            pdf_data
        ).decode('utf-8')

        nylas_url = (
            'https://api.us.nylas.com/v3/grants/'
            f'{nylas_grant_id}/messages/send'
        )

        email_data = {
            'to': [
                {
                    'name': 'Muthyala Tejas',
                    'email': 'muthyalatejas8055@gmail.com',
                }
            ],
            'subject': (
                f'MediCare Plus - Patient Medical Report '
                f'{patient.patient_id}'
            ),
            'body': (
                '<h2>MediCare Plus Hospital</h2>'
                '<p>Dear Tejas,</p>'
                f'<p>Please find attached the requested '
                f'<strong>Patient Medical Report</strong> '
                f'for patient <strong>{patient.patient_id}</strong>.</p>'
                '<p>The report was generated from the '
                'MediCare Plus Hospital Management System.</p>'
                '<p>Regards,<br>'
                '<strong>MediCare Plus Hospital</strong><br>'
                'Hospital Management System</p>'
            ),
            'attachments': [
                {
                    'filename': filename,
                    'content': encoded_pdf,
                    'content_type': 'application/pdf',
                }
            ],
        }

        nylas_response = requests.post(
            nylas_url,
            headers={
                'Authorization': f'Bearer {nylas_api_key}',
                'Content-Type': 'application/json',
            },
            json=email_data,
            timeout=30,
        )

        print(
            'Nylas email status:',
            nylas_response.status_code
        )
        print(
            'Nylas email response:',
            nylas_response.text
        )

        if not nylas_response.ok:
            raise ValueError(
                'Nylas email request failed.'
            )

        messages.success(
            request,
            f'PDF downloaded and emailed successfully for '
            f'{patient.patient_id}.'
        )

    except Exception as error:
        print(
            'Nylas email error:',
            error
        )

        messages.warning(
            request,
            'PDF was generated and downloaded, but the email '
            'could not be sent.'
        )

    # Download the same PDF in the browser
    response = HttpResponse(
        pdf_data,
        content_type='application/pdf'
    )

    response['Content-Disposition'] = (
        f'attachment; filename="{filename}"'
    )

    return response


@login_required
def dashboard_download_excel(request):
    """
    Download all patient records as an Excel file.
    """

    patients = Patient.objects.order_by(
        'patient_id'
    )

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = 'Patients'

    headers = [
        'Patient ID',
        'Date of Birth',
        'Gender',
        'Phone',
        'Address',
        'Blood Group',
        'Emergency Contact',
    ]

    worksheet.append(headers)

    for patient in patients:

        worksheet.append([
            patient.patient_id,
            patient.date_of_birth,
            patient.get_gender_display(),
            patient.phone,
            patient.address,
            patient.blood_group or '',
            patient.emergency_contact or '',
        ])

    column_widths = {
        'A': 15,
        'B': 18,
        'C': 15,
        'D': 15,
        'E': 35,
        'F': 15,
        'G': 20,
    }

    for column, width in column_widths.items():

        worksheet.column_dimensions[
            column
        ].width = width

    response = HttpResponse(
        content_type=(
            'application/vnd.openxmlformats-officedocument.'
            'spreadsheetml.sheet'
        )
    )

    response['Content-Disposition'] = (
        'attachment; '
        'filename="medicare_plus_patients.xlsx"'
    )

    workbook.save(response)

    return response